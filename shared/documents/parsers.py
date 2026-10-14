import os
import uuid
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Optional, Any
import mimetypes

from .models import Document, DocumentChunk, DocumentMetadata

class DocumentParser(ABC):
    @abstractmethod
    def parse(self, filepath: Path) -> Document:
        pass
    
    def _create_base_doc(self, filepath: Path) -> Document:
        mime_type, _ = mimetypes.guess_type(str(filepath))
        return Document(
            doc_id=str(uuid.uuid4()),
            filename=filepath.name,
            mime_type=mime_type or "application/octet-stream",
            size_bytes=filepath.stat().st_size,
            source_path=filepath
        )

class TextParser(DocumentParser):
    def parse(self, filepath: Path) -> Document:
        doc = self._create_base_doc(filepath)
        content = ""
        for encoding in ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']:
            try:
                with open(filepath, 'r', encoding=encoding) as f:
                    content = f.read()
                break
            except UnicodeDecodeError:
                continue
                
        paragraphs = content.split('\n\n')
        for i, para in enumerate(paragraphs):
            para = para.strip()
            if para:
                doc.chunks.append(DocumentChunk(
                    chunk_id=f"{doc.doc_id}_chunk_{i}",
                    content=para,
                    source_reference=f"Paragraph {i+1}"
                ))
        return doc

class PDFParser(DocumentParser):
    def parse(self, filepath: Path) -> Document:
        import fitz
        doc = self._create_base_doc(filepath)
        try:
            pdf_doc = fitz.open(filepath)
            meta = pdf_doc.metadata
            if meta:
                doc.metadata.title = meta.get("title", "")
                doc.metadata.author = meta.get("author", "")
                doc.metadata.creation_date = meta.get("creationDate", "")
                
            for page_num, page in enumerate(pdf_doc):
                text = page.get_text()
                if text.strip():
                    doc.chunks.append(DocumentChunk(
                        chunk_id=f"{doc.doc_id}_p{page_num+1}",
                        content=text.strip(),
                        source_reference=f"Page {page_num+1}"
                    ))
            pdf_doc.close()
        except Exception as e:
            raise ValueError(f"Failed to parse PDF: {e}")
        return doc

class DocxParser(DocumentParser):
    def parse(self, filepath: Path) -> Document:
        import docx
        doc = self._create_base_doc(filepath)
        try:
            docx_doc = docx.Document(filepath)
            core_props = docx_doc.core_properties
            doc.metadata.title = core_props.title or ""
            doc.metadata.author = core_props.author or ""
            
            current_chunk = []
            chunk_idx = 1
            for para in docx_doc.paragraphs:
                if para.text.strip():
                    current_chunk.append(para.text.strip())
                if len(current_chunk) >= 5:
                    doc.chunks.append(DocumentChunk(
                        chunk_id=f"{doc.doc_id}_chunk_{chunk_idx}",
                        content="\n".join(current_chunk),
                        source_reference=f"Section {chunk_idx}"
                    ))
                    current_chunk = []
                    chunk_idx += 1
            if current_chunk:
                doc.chunks.append(DocumentChunk(
                    chunk_id=f"{doc.doc_id}_chunk_{chunk_idx}",
                    content="\n".join(current_chunk),
                    source_reference=f"Section {chunk_idx}"
                ))
            for table_idx, table in enumerate(docx_doc.tables):
                table_text = []
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells]
                    table_text.append(" | ".join(row_text))
                if table_text:
                    doc.chunks.append(DocumentChunk(
                        chunk_id=f"{doc.doc_id}_table_{table_idx+1}",
                        content="\n".join(table_text),
                        source_reference=f"Table {table_idx+1}"
                    ))
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {e}")
        return doc

class PptxParser(DocumentParser):
    def parse(self, filepath: Path) -> Document:
        import pptx
        doc = self._create_base_doc(filepath)
        try:
            prs = pptx.Presentation(filepath)
            core_props = prs.core_properties
            doc.metadata.title = core_props.title or ""
            doc.metadata.author = core_props.author or ""
            
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text.strip())
                if slide_text:
                    doc.chunks.append(DocumentChunk(
                        chunk_id=f"{doc.doc_id}_slide_{i+1}",
                        content="\n".join(slide_text),
                        source_reference=f"Slide {i+1}"
                    ))
        except Exception as e:
            raise ValueError(f"Failed to parse PPTX: {e}")
        return doc

class PandasParser(DocumentParser):
    def parse(self, filepath: Path) -> Document:
        import pandas as pd
        doc = self._create_base_doc(filepath)
        ext = filepath.suffix.lower()
        try:
            if ext in ['.xlsx', '.xls']:
                with pd.ExcelFile(filepath) as excel_file:
                    for sheet_name in excel_file.sheet_names:
                        df = pd.read_excel(excel_file, sheet_name=sheet_name)
                        markdown_table = df.to_markdown(index=False)
                        if markdown_table:
                            doc.chunks.append(DocumentChunk(
                                chunk_id=f"{doc.doc_id}_sheet_{sheet_name}",
                                content=markdown_table,
                                source_reference=f"Sheet: {sheet_name}"
                            ))
            elif ext == '.csv':
                df = pd.read_csv(filepath)
                markdown_table = df.to_markdown(index=False)
                if markdown_table:
                    doc.chunks.append(DocumentChunk(
                        chunk_id=f"{doc.doc_id}_csv",
                        content=markdown_table,
                        source_reference="CSV Data"
                    ))
        except Exception as e:
            raise ValueError(f"Failed to parse Spreadsheet: {e}")
        return doc

class ImageParser(DocumentParser):
    def parse(self, filepath: Path) -> Document:
        doc = self._create_base_doc(filepath)
        content_lines = [f"[Image File: {filepath.name}]"]
        
        try:
            import cv2
            from jarvis_ai.vision.object_detection import ObjectDetector
            
            detector = ObjectDetector()
            frame = cv2.imread(str(filepath))
            if frame is not None:
                h, w = frame.shape[:2]
                content_lines.append(f"Dimensions: {w}x{h}")
                
                # Run YOLO
                detections = detector.detect(frame)
                if detections:
                    counts = {}
                    for d in detections:
                        name = d.get('name', 'object')
                        counts[name] = counts.get(name, 0) + 1
                    
                    found_objs = [f"{count} {name}(s)" for name, count in counts.items()]
                    content_lines.append("Detected Objects (YOLOv8): " + ", ".join(found_objs))
                else:
                    content_lines.append("Detected Objects: None detected.")
            else:
                content_lines.append("Warning: Could not read image frame.")
        except Exception as e:
            content_lines.append(f"Vision Analysis Failed: {e}")
            
        content_lines.append("Note: If the user asks for OCR or more detailed analysis, prompt them to use a dedicated vision model.")
        
        doc.chunks.append(DocumentChunk(
            chunk_id=f"{doc.doc_id}_img",
            content="\n".join(content_lines),
            source_reference="Image Analysis"
        ))
        return doc

class ParserRegistry:
    def __init__(self):
        self.parsers = {
            ".txt": TextParser(),
            ".md": TextParser(),
            ".json": TextParser(),
            ".py": TextParser(),
            ".js": TextParser(),
            ".ts": TextParser(),
            ".java": TextParser(),
            ".cpp": TextParser(),
            ".c": TextParser(),
            ".cs": TextParser(),
            ".html": TextParser(),
            ".css": TextParser(),
            ".sql": TextParser(),
            ".ps1": TextParser(),
            ".pdf": PDFParser(),
            ".docx": DocxParser(),
            ".pptx": PptxParser(),
            ".xlsx": PandasParser(),
            ".csv": PandasParser(),
            ".png": ImageParser(),
            ".jpg": ImageParser(),
            ".jpeg": ImageParser(),
            ".gif": ImageParser()
        }
        
    def parse_file(self, filepath: Path) -> Document:
        ext = filepath.suffix.lower()
        parser = self.parsers.get(ext)
        if not parser:
            parser = TextParser()
        return parser.parse(filepath)
