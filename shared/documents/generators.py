import os
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Optional
from shared.config.path_manager import PathManager

class FileGenerator(ABC):
    @abstractmethod
    def generate(self, content: str, output_path: Path) -> Path:
        pass

class TextGenerator(FileGenerator):
    def generate(self, content: str, output_path: Path) -> Path:
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(content)
        return output_path

class PdfGenerator(FileGenerator):
    def generate(self, content: str, output_path: Path) -> Path:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        doc = SimpleDocTemplate(str(output_path), pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        for para in content.split('\n\n'):
            if para.strip():
                story.append(Paragraph(para.strip(), styles["Normal"]))
                story.append(Spacer(1, 12))
        doc.build(story)
        return output_path

class DocxGenerator(FileGenerator):
    def generate(self, content: str, output_path: Path) -> Path:
        import docx
        doc = docx.Document()
        for para in content.split('\n\n'):
            if para.strip():
                doc.add_paragraph(para.strip())
        doc.save(output_path)
        return output_path

class PptxGenerator(FileGenerator):
    def generate(self, content: str, output_path: Path) -> Path:
        import pptx
        prs = pptx.Presentation()
        paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]
        title_slide_layout = prs.slide_layouts[0]
        bullet_slide_layout = prs.slide_layouts[1]
        if paragraphs:
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            title.text = paragraphs[0][:50] + "..." if len(paragraphs[0]) > 50 else paragraphs[0]
            if len(paragraphs) > 1:
                slide.placeholders[1].text = paragraphs[1][:100]
        for para in paragraphs[1:]:
            slide = prs.slides.add_slide(bullet_slide_layout)
            shapes = slide.shapes
            title_shape = shapes.title
            body_shape = shapes.placeholders[1]
            title_shape.text = "Generated Content"
            tf = body_shape.text_frame
            tf.text = para
        prs.save(output_path)
        return output_path

class XlsxGenerator(FileGenerator):
    def generate(self, content: str, output_path: Path) -> Path:
        import pandas as pd
        import io
        import csv
        
        # Try to parse content as CSV first to make a structured Excel
        try:
            # If it looks like CSV
            if ',' in content and '\n' in content:
                df = pd.read_csv(io.StringIO(content))
            else:
                df = pd.DataFrame([line.split('\t') for line in content.split('\n') if line.strip()])
        except Exception:
            # Fallback to single column
            df = pd.DataFrame({"Content": [line for line in content.split('\n') if line.strip()]})
            
        df.to_excel(str(output_path), index=False)
        return output_path

class CsvGenerator(FileGenerator):
    def generate(self, content: str, output_path: Path) -> Path:
        # Just write it out, assuming the LLM formatted it as CSV.
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(content)
        return output_path

class GeneratorManager:
    def __init__(self):
        self.generators = {
            ".txt": TextGenerator(),
            ".md": TextGenerator(),
            ".json": TextGenerator(),
                        ".csv": CsvGenerator(),
            ".xlsx": XlsxGenerator(),
            ".pdf": PdfGenerator(),
            ".docx": DocxGenerator(),
            ".pptx": PptxGenerator(),
        }
        
    def _get_safe_output_dir(self) -> Path:
        docs_dir = Path(os.path.expanduser("~/Documents"))
        gen_dir = docs_dir / "AI Suite" / "Generated"
        gen_dir.mkdir(parents=True, exist_ok=True)
        return gen_dir
        
    def _get_safe_filename(self, base_name: str, ext: str) -> Path:
        import re
        # Sanitize against path traversal and invalid characters
        base_name = Path(base_name).name
        base_name = re.sub(r'[\/*?:"<>|]', "", base_name)
        if not base_name.strip() or base_name in (".", ".."):
            base_name = "generated_file"
            
        gen_dir = self._get_safe_output_dir()
        filename = f"{base_name}{ext}"
        filepath = gen_dir / filename
        counter = 1
        while filepath.exists():
            filepath = gen_dir / f"{base_name}_{counter}{ext}"
            counter += 1
        return filepath
        
    def generate_file(self, content: str, base_name: str, ext: str) -> Path:
        ext = ext.lower()
        if not ext.startswith("."):
            ext = "." + ext
        generator = self.generators.get(ext)
        if not generator:
            generator = TextGenerator()
            ext = ".txt"
        filepath = self._get_safe_filename(base_name, ext)
        return generator.generate(content, filepath)
