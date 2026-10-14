import os
import unittest
from pathlib import Path
import json
import csv
import pandas as pd
import fitz
import docx
import pptx

from shared.documents.generators import GeneratorManager
from shared.documents.parsers import ParserRegistry
from shared.documents.models import Document

class TestDeepValidation(unittest.TestCase):
    
    def setUp(self):
        self.test_dir = Path("tests/test_deep_gen")
        self.test_dir.mkdir(parents=True, exist_ok=True)
        self.gen = GeneratorManager()
        self.gen._get_safe_output_dir = lambda: self.test_dir
        self.registry = ParserRegistry()

    def tearDown(self):
        import shutil
        if self.test_dir.exists():
            shutil.rmtree(self.test_dir)

    def _generate_and_check(self, content: str, ext: str, expected_content: str = None) -> Path:
        out_path = self.gen.generate_file(content, "test_file", ext)
        self.assertTrue(out_path.exists())
        self.assertEqual(out_path.suffix.lower(), ext)
        return out_path
        
    def test_txt_generation_and_parsing(self):
        p = self._generate_and_check("Hello Text", ".txt")
        self.assertEqual(p.read_text(encoding="utf-8"), "Hello Text")
        
        doc = self.registry.parse_file(p)
        self.assertIn("Hello Text", doc.full_text)

    def test_json_generation_and_parsing(self):
        data = {"key": "val", "list": [1, 2, 3]}
        content = json.dumps(data)
        p = self._generate_and_check(content, ".json")
        
        with open(p, "r", encoding="utf-8") as f:
            parsed = json.load(f)
            self.assertEqual(parsed["key"], "val")
            
        doc = self.registry.parse_file(p)
        self.assertIn('"val"', doc.full_text)

    def test_csv_generation_and_parsing(self):
        content = "Name,Age\nAlice,30\nBob,25"
        p = self._generate_and_check(content, ".csv")
        
        with open(p, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)
            self.assertEqual(rows[0], ["Name", "Age"])
            self.assertEqual(rows[1], ["Alice", "30"])
            
        doc = self.registry.parse_file(p)
        self.assertIn("Alice", doc.full_text)

    def test_xlsx_generation_and_parsing(self):
        content = "Name,Age\nAlice,30\nBob,25"
        p = self._generate_and_check(content, ".xlsx")
        
        df = pd.read_excel(p)
        self.assertEqual(list(df.columns), ["Name", "Age"])
        self.assertEqual(str(df.iloc[0]["Name"]), "Alice")
        
        doc = self.registry.parse_file(p)
        self.assertIn("Alice", doc.full_text)

    def test_docx_generation_and_parsing(self):
        content = "First Paragraph\n\nSecond Paragraph"
        p = self._generate_and_check(content, ".docx")
        
        d = docx.Document(p)
        paras = [p.text for p in d.paragraphs if p.text.strip()]
        self.assertEqual(paras[0], "First Paragraph")
        self.assertEqual(paras[1], "Second Paragraph")
        
        doc = self.registry.parse_file(p)
        self.assertIn("First Paragraph", doc.full_text)
        self.assertIn("Second Paragraph", doc.full_text)

    def test_pptx_generation_and_parsing(self):
        content = "Title Slide\n\nBullet 1\n\nBullet 2"
        p = self._generate_and_check(content, ".pptx")
        
        prs = pptx.Presentation(p)
        slides = prs.slides
        self.assertEqual(len(slides), 3)
        self.assertEqual(slides[0].shapes.title.text, "Title Slide")
        self.assertEqual(slides[1].shapes.placeholders[1].text_frame.text, "Bullet 1")
        
        doc = self.registry.parse_file(p)
        self.assertIn("Title Slide", doc.full_text)
        self.assertIn("Bullet 1", doc.full_text)

    def test_pdf_generation_and_parsing(self):
        content = "PDF Paragraph 1\n\nPDF Paragraph 2"
        p = self._generate_and_check(content, ".pdf")
        
        doc_pdf = fitz.open(p)
        text = ""
        for page in doc_pdf:
            text += page.get_text()
        self.assertIn("PDF Paragraph 1", text)
        self.assertIn("PDF Paragraph 2", text)
        doc_pdf.close()
        
        doc = self.registry.parse_file(p)
        self.assertIn("PDF Paragraph 1", doc.full_text)
        self.assertIn("PDF Paragraph 2", doc.full_text)

if __name__ == "__main__":
    unittest.main()
